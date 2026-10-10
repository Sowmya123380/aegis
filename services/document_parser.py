"""Aegis Document Intelligence Agent - Document Parser Service.

Handles document parsing for PDF, DOCX, and CSV formats with:
- Per-page text extraction for PDFs
- Paragraph and formatted table extraction for DOCX
- Robust CSV reading with multiple encoding fallbacks
- Scanned PDF detection and OCR notice generation
- Malformed file and exception safety
"""

from io import BytesIO
from typing import Any, Dict, List, Optional, Union
import pandas as pd
import pymupdf as fitz
from docx import Document


def parse_pdf(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse PDF file with per-page text tracking and scanned document detection."""
    pages: List[Dict[str, Any]] = []
    full_text_list: List[str] = []

    try:
        with fitz.open(stream=raw_bytes, filetype="pdf") as pdf:
            page_count = len(pdf)
            for page_idx, page in enumerate(pdf, start=1):
                page_text = page.get_text("text").strip()
                pages.append({
                    "page_number": page_idx,
                    "text": page_text,
                    "character_count": len(page_text),
                })
                if page_text:
                    full_text_list.append(f"--- Page {page_idx} ---\n{page_text}")

            total_text = "\n\n".join(full_text_list).strip()
            total_chars = len(total_text)
            
            # Scanned document detection heuristic
            is_scanned = page_count > 0 and (total_chars / page_count < 25)
            if is_scanned:
                notice = f"[NOTICE: '{filename}' appears to be a scanned or image-based PDF. Extracted text is sparse ({total_chars} chars across {page_count} pages). OCR layer recommended.]"
                total_text = f"{notice}\n\n{total_text}".strip()

            return {
                "name": filename,
                "type": "pdf",
                "text": total_text,
                "content": total_text,
                "pages": pages,
                "page_count": page_count,
                "character_count": total_chars,
                "is_scanned": is_scanned,
                "dataframe": None,
                "df": None,
                "metadata": {
                    "format": "PDF",
                    "page_count": page_count,
                    "is_scanned": is_scanned,
                }
            }
    except Exception as exc:
        raise ValueError(f"Failed to parse PDF file '{filename}': {type(exc).__name__} - {exc}")


def parse_docx(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse DOCX file with paragraph and table extraction."""
    try:
        doc = Document(BytesIO(raw_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        table_lines: List[str] = []
        for table_idx, table in enumerate(doc.tables, start=1):
            table_lines.append(f"--- Table {table_idx} ---")
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(row_cells))

        combined_lines = paragraphs + table_lines
        total_text = "\n".join(combined_lines).strip()

        pages = [
            {"page_number": 1, "text": total_text, "character_count": len(total_text)}
        ]

        return {
            "name": filename,
            "type": "docx",
            "text": total_text,
            "content": total_text,
            "pages": pages,
            "page_count": 1,
            "character_count": len(total_text),
            "is_scanned": False,
            "dataframe": None,
            "df": None,
            "metadata": {
                "format": "DOCX",
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
            }
        }
    except Exception as exc:
        raise ValueError(f"Failed to parse DOCX file '{filename}': {type(exc).__name__} - {exc}")


def parse_csv(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse CSV file with encoding fallback and pandas dataframe creation."""
    encodings = ["utf-8", "latin1", "cp1252", "iso-8859-1"]
    df: Optional[pd.DataFrame] = None
    last_err: Optional[Exception] = None

    for enc in encodings:
        try:
            df = pd.read_csv(BytesIO(raw_bytes), encoding=enc)
            break
        except Exception as exc:
            last_err = exc

    if df is None:
        raise ValueError(f"Failed to parse CSV file '{filename}': {last_err}")

    text_preview = df.head(100).to_string(index=False)
    pages = [
        {"page_number": 1, "text": text_preview, "character_count": len(text_preview)}
    ]

    return {
        "name": filename,
        "type": "csv",
        "text": text_preview,
        "content": text_preview,
        "pages": pages,
        "page_count": 1,
        "character_count": len(text_preview),
        "is_scanned": False,
        "dataframe": df,
        "df": df,
        "metadata": {
            "format": "CSV",
            "rows": len(df),
            "columns": len(df.columns),
            "column_names": [str(c) for c in df.columns],
        }
    }


def parse_document(upload_or_bytes: Union[Any, bytes], filename: Optional[str] = None) -> Dict[str, Any]:
    """Unified entrypoint for parsing uploaded files or byte streams into structured documents."""
    if hasattr(upload_or_bytes, "name") and hasattr(upload_or_bytes, "getvalue"):
        file_name = upload_or_bytes.name
        raw_bytes = upload_or_bytes.getvalue()
    elif isinstance(upload_or_bytes, bytes) and filename:
        file_name = filename
        raw_bytes = upload_or_bytes
    else:
        raise ValueError("Invalid document input. Must provide uploaded file object or raw bytes with filename.")

    ext = file_name.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        return parse_pdf(raw_bytes, file_name)
    elif ext == "docx":
        return parse_docx(raw_bytes, file_name)
    elif ext == "csv":
        return parse_csv(raw_bytes, file_name)
    else:
        raise ValueError(f"Unsupported document format '.{ext}' for file '{file_name}'. Supported: PDF, DOCX, CSV.")
